"""Vehicle makes and models (owner's request N8, migration 1157): a list reviewed by a second person, carriers' requests
for a missing model, the owner's table imported for review, and vehicles that pick their model."""
import io
import uuid

import pytest

from test_e2e import OWNER_URL, login, owner_sql
from test_users_and_roles import NEW_PASSWORD, add_staff, role_uid

pytestmark = pytest.mark.skipif(not OWNER_URL, reason="needs MASSLAK_OWNER_URL")
AR = "\u0645\u0631\u0643\u0628\u0629"


@pytest.fixture(scope="module")
def admin():
    return login("admin@masslak.test", "PLATFORM")


@pytest.fixture(scope="module")
def second(admin):
    email, _ = add_staff(admin, [role_uid(admin, "PLATFORM_OPERATIONS")])
    return login(email, "PLATFORM", NEW_PASSWORD)


@pytest.fixture(scope="module")
def carrier():
    return login("owner@carrier.test", "OPERATOR")


def tag() -> str:
    return uuid.uuid4().hex[:6].upper()


def accepted_model(admin, second) -> tuple[str, str, str]:
    name = "Make" + tag()
    make = admin.post("/api/admin/vehicle-models/makes", json={"name_en": name, "name_ar": AR, "country_code": "DE"}).json()["uid"]
    model = admin.post("/api/admin/vehicle-models/models",
                       json={"make_uid": make, "name": "T" + tag(), "vehicle_type": "COACH", "seats_from": 45, "seats_to": 53,
                             "confidence": "HIGH"})
    assert model.status_code == 201, model.text
    model = model.json()["uid"]
    r = second.post(f"/api/admin/vehicle-models/models/{model}/review", json={"decision": "ACCEPT"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "MAKE_NOT_ACTIVE"          # the make first
    assert second.post(f"/api/admin/vehicle-models/makes/{make}/review", json={"decision": "ACCEPT"}).json()["status"] == "ACTIVE"
    assert second.post(f"/api/admin/vehicle-models/models/{model}/review", json={"decision": "ACCEPT"}).json()["status"] == "ACTIVE"
    return make, model, name


def carrier_models(carrier) -> dict:
    return {m["uid"]: m for k in carrier.get("/api/vehicle-models").json()["makes"] for m in k["models"]}


def test_a_model_reaches_carriers_after_a_second_review_and_names_the_vehicle(admin, second, carrier):
    name = "Make" + tag()
    make = admin.post("/api/admin/vehicle-models/makes", json={"name_en": name, "name_ar": AR}).json()["uid"]
    own = admin.post(f"/api/admin/vehicle-models/makes/{make}/review", json={"decision": "ACCEPT"})
    assert own.status_code == 409 and own.json()["error"]["code"] == "FOUR_EYES"
    make, model, make_name = accepted_model(admin, second)
    seen = carrier_models(carrier)[model]
    assert seen["seats_from"] == 45 and "notes" not in seen
    plate = tag()
    layout = carrier.get("/api/carrier/seat-layouts").json()["layouts"][0]["uid"]
    r = carrier.post("/api/carrier/vehicles", json={
        "plate_no": f"MM-{plate}", "chassis_no": f"CHMM{plate}", "vehicle_type": "COACH", "model_uid": model, "seat_layout_uid": layout,
        "insurance_no": f"POL-M{plate}", "insurer": "Test", "insurance_issue": "2026-01-01", "insurance_expiry": "2030-01-01"})
    assert r.status_code == 201, r.text
    row = owner_sql("SELECT make || '|' || model FROM fleet.vehicle WHERE uid = $1", uuid.UUID(r.json()["uid"]))
    assert row.startswith(make_name + "|")                          # names copied from the list
    # a model still waiting for review cannot be picked
    pending = admin.post("/api/admin/vehicle-models/models", json={"make_uid": make, "name": "P" + tag(), "vehicle_type": "VAN"}).json()["uid"]
    r = carrier.post("/api/carrier/vehicles", json={
        "plate_no": f"MP-{plate}", "chassis_no": f"CHMP{plate}", "vehicle_type": "VAN", "model_uid": pending, "seat_layout_uid": layout,
        "insurance_no": f"POL-P{plate}", "insurer": "Test", "insurance_issue": "2026-01-01", "insurance_expiry": "2030-01-01"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "MODEL_UNKNOWN"
    assert carrier.post("/api/admin/vehicle-models/makes", json={"name_en": "X" + tag(), "name_ar": AR}).status_code == 403


def test_a_carrier_asks_for_a_missing_model_and_the_platform_decides(admin, second, carrier):
    r = carrier.post("/api/vehicle-models/requests", json={"make_name": "Newmake" + tag(), "model_name": "NM1", "vehicle_type": "MINIBUS",
                                                          "seats": 28, "note": "two in our fleet"})
    assert r.status_code == 201, r.text
    req = r.json()["uid"]
    mine = {x["uid"]: x for x in carrier.get("/api/vehicle-models/requests").json()["requests"]}
    assert mine[req]["status"] == "SUBMITTED"
    r = admin.post(f"/api/admin/vehicle-models/requests/{req}/decide", json={"decision": "APPROVE"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "MAKE_NAME_AR"
    assert admin.post(f"/api/admin/vehicle-models/requests/{req}/decide",
                      json={"decision": "APPROVE", "make_name_ar": AR}).json()["status"] == "APPROVED"
    # the new make and model still wait for a second person
    status = owner_sql("""SELECT m.status || '|' || k.status || '|' || m.source FROM fleet.vehicle_model_request r
                           JOIN ref.vehicle_model m ON m.id = r.model_id JOIN ref.vehicle_make k ON k.id = m.make_id WHERE r.uid = $1""",
                       uuid.UUID(req))
    assert status == "PENDING_REVIEW|PENDING_REVIEW|REQUEST"
    again = admin.post(f"/api/admin/vehicle-models/requests/{req}/decide", json={"decision": "REJECT", "note": "late"})
    assert again.json()["error"]["code"] == "REQUEST_HANDLED"

    _, model, _ = accepted_model(admin, second)
    r2 = carrier.post("/api/vehicle-models/requests", json={"make_name": "Dup" + tag(), "model_name": "D1", "vehicle_type": "COACH"}).json()["uid"]
    assert admin.post(f"/api/admin/vehicle-models/requests/{r2}/decide", json={"decision": "MATCH", "model_uid": model}).json()["status"] == "MATCHED"
    r3 = carrier.post("/api/vehicle-models/requests", json={"make_name": "Nope" + tag(), "model_name": "N1", "vehicle_type": "VAN"}).json()["uid"]
    assert admin.post(f"/api/admin/vehicle-models/requests/{r3}/decide", json={"decision": "REJECT"}).json()["error"]["code"] == "REASON_REQUIRED"
    assert admin.post(f"/api/admin/vehicle-models/requests/{r3}/decide",
                      json={"decision": "REJECT", "note": "not a passenger vehicle"}).json()["status"] == "REJECTED"
    r4 = carrier.post("/api/vehicle-models/requests", json={"make_name": "W" + tag(), "model_name": "W1", "vehicle_type": "VAN"}).json()["uid"]
    assert carrier.post(f"/api/vehicle-models/requests/{r4}/withdraw").json()["status"] == "WITHDRAWN"


def _workbook(rows: list[list]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "vehicle_models"
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_the_owner_table_comes_in_waiting_for_review(admin):
    head = ["#", "make_en", "make_ar", "country", "model", "vehicle_class", "vehicle_type", "seats_from (approx.)", "seats_to (approx.)",
            "truck_weight_class", "status", "confidence", "notes", "owner_decision (keep/remove/edit)"]
    make = "Imp" + tag()
    rows = [head,
            [1, make, AR, "China", "C10", "BUS", "COACH", 49, 53, None, "current", "verify", "name to confirm", None],
            [2, make, AR, "China", "T20", "TRUCK", None, None, None, "heavy", "older", "high", None, None],
            [3, make, AR, "China", "Gone", "BUS", "COACH", 30, 30, None, "current", "high", None, "remove"]]
    bad = admin.post("/api/admin/vehicle-models/import", files={"file": ("m.xlsx", _workbook(rows + [[4, make, AR, "China", "Z", "BOAT", None, None, None,
                                                                                                    None, "current", "high", None, None]]))})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "REGISTRY_FILE_ROWS", bad.text
    r = admin.post("/api/admin/vehicle-models/import", files={"file": ("m.xlsx", _workbook(rows))})
    assert r.status_code == 200 and r.json() == {"makes": 1, "models": 2, "trims": 0, "skipped": 0}, r.text
    assert admin.post("/api/admin/vehicle-models/import", files={"file": ("m.xlsx", _workbook(rows))}).json()["skipped"] == 2
    found = next(k for k in admin.get("/api/admin/vehicle-models").json()["makes"] if k["name_en"] == make)
    assert found["status"] == "PENDING_REVIEW" and found["country_code"] == "CN" and found["source"] == "IMPORT"
    truck = next(m for m in found["models"] if m["name"] == "T20")
    assert truck["vehicle_type"] == "OTHER" and truck["weight_class"] == "HEAVY" and truck["production"] == "OLDER"
    assert {m["name"] for m in found["models"]} == {"C10", "T20"}                  # the row marked "remove" stays out


def _book(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_the_2020_list_comes_with_arabic_names_years_and_trims(admin):
    make = "Yr" + tag()
    models = [["make_en", "make_ar", "country", "model", "model_ar", "vehicle_class", "vehicle_type", "body_type", "seats_from",
               "seats_to", "year_from", "year_to", "status", "confidence"],
              [make, AR, "Japan", "Saloon", AR, "CAR", "OTHER", "SEDAN", 5, 5, 2020, None, "current", "high"],
              [make, AR, "Japan", "Old", AR, "CAR", "OTHER", "HATCH", 5, 5, 2017, 2021, "older", "high"]]
    trims = [["make_en", "model", "trim", "trim_ar", "fuel", "year_from", "year_to"],
             [make, "Saloon", "LE", AR, "Petrol", 2020, None], [make, "Saloon", "LE Hybrid", AR, "Hybrid", 2020, None],
             [make, "Nowhere", "X", AR, "Petrol", 2020, None]]
    r = admin.post("/api/admin/vehicle-models/import", files={"file": ("m.xlsx", _book({"vehicle_models_import": models, "trims": trims}))})
    assert r.status_code == 200 and r.json() == {"makes": 1, "models": 2, "trims": 2, "skipped": 1}, r.text
    bad = trims + [[make, "Saloon", "Z", AR, "Steam", 2020, None]]
    r = admin.post("/api/admin/vehicle-models/import", files={"file": ("m.xlsx", _book({"vehicle_models_import": models, "trims": bad}))})
    assert r.status_code == 422 and "fuel" in r.json()["error"]["message"]
    found = next(k for k in admin.get("/api/admin/vehicle-models").json()["makes"] if k["name_en"] == make)
    saloon = next(m for m in found["models"] if m["name"] == "Saloon")
    assert saloon["name_ar"] == AR and saloon["body_type"] == "SEDAN" and saloon["year_from"] == 2020 and saloon["year_to"] is None
    assert {(t["name"], t["fuel"]) for t in saloon["trims"]} == {("LE", "PETROL"), ("LE Hybrid", "HYBRID")}
    assert next(m for m in found["models"] if m["name"] == "Old")["production"] == "OLDER"


def test_entries_are_corrected_with_history_and_deleted_only_while_unused(admin, second, carrier):
    make, model, name = accepted_model(admin, second)
    r = admin.patch(f"/api/admin/vehicle-models/models/{model}", json={"name": "Renamed " + tag()})
    assert r.status_code == 422 and r.json()["error"]["code"] == "MODEL_CHANGE_REASON"
    assert admin.patch(f"/api/admin/vehicle-models/models/{model}", json={"year_from": 2024, "year_to": 2021}).json()["error"]["code"] == "YEARS_RANGE"
    r = admin.patch(f"/api/admin/vehicle-models/models/{model}", json={"name_ar": AR, "body_type": "COACH", "year_from": 2020,
                                                                       "seats_to": 57, "reason": "the maker's catalogue"})
    assert r.status_code == 200 and set(r.json()["changed"]) == {"name_ar", "body_type", "year_from", "seats_to"}, r.text
    assert admin.patch(f"/api/admin/vehicle-models/makes/{make}", json={"name_ar": AR + AR}).json()["error"]["code"] == "MAKE_CHANGE_REASON"
    assert admin.patch(f"/api/admin/vehicle-models/makes/{make}", json={"name_ar": AR + AR, "reason": "spelt as the dealer"}).status_code == 200
    t = admin.post(f"/api/admin/vehicle-models/models/{model}/trims", json={"name": "Lux", "fuel": "DIESEL", "year_from": 2021})
    assert t.status_code == 201, t.text
    trim = t.json()["uid"]
    assert admin.post(f"/api/admin/vehicle-models/models/{model}/trims", json={"name": "lux", "fuel": "DIESEL"}).json()["error"]["code"] == "TRIM_EXISTS"
    assert admin.patch(f"/api/admin/vehicle-models/trims/{trim}", json={"fuel": "HYBRID"}).status_code == 200
    d = admin.get(f"/api/admin/vehicle-models/models/{model}").json()
    assert d["model"]["seats_to"] == 57 and d["used_by"] == 0
    assert {h["field"] for h in d["history"]} >= {"name_ar", "body_type", "year_from", "seats_to"}
    assert next(h for h in d["history"] if h["field"] == "seats_to")["reason"] == "the maker's catalogue"
    assert {h["field"] for h in admin.get(f"/api/admin/vehicle-models/trims/{trim}").json()["history"]} == {"fuel"}
    assert carrier.patch(f"/api/admin/vehicle-models/models/{model}", json={"notes": "x"}).status_code in (401, 403)
    # a model a vehicle uses is not deleted; nor its make while it has models
    layouts = carrier.get("/api/carrier/seat-layouts").json()["layouts"]
    v = carrier.post("/api/carrier/vehicles", json={"plate_no": "DL-" + tag(), "chassis_no": "CHS" + tag() + "DEL", "vehicle_type": "COACH",
                                                     "seat_layout_uid": layouts[0]["uid"], "model_uid": model, "insurance_no": "INS-" + tag(),
                                                     "insurer": "Test Insurance", "insurance_issue": "2026-01-01", "insurance_expiry": "2027-12-31"})
    assert v.status_code == 201, v.text
    r = admin.delete(f"/api/admin/vehicle-models/models/{model}")
    assert r.status_code == 409 and r.json()["error"]["code"] == "MODEL_IN_USE"
    assert admin.delete(f"/api/admin/vehicle-models/makes/{make}").json()["error"]["code"] == "MAKE_HAS_MODELS"
    # an unused model goes, with its trims and history; then its make
    name2 = "Del" + tag()
    mk2 = admin.post("/api/admin/vehicle-models/makes", json={"name_en": name2, "name_ar": AR}).json()["uid"]
    m2 = admin.post("/api/admin/vehicle-models/models", json={"make_uid": mk2, "name": "Gone", "vehicle_type": "VAN"}).json()["uid"]
    admin.post(f"/api/admin/vehicle-models/models/{m2}/trims", json={"name": "Base"})
    admin.patch(f"/api/admin/vehicle-models/models/{m2}", json={"notes": "entered by mistake"})
    assert admin.delete(f"/api/admin/vehicle-models/models/{m2}").json()["deleted"] is True
    assert admin.get(f"/api/admin/vehicle-models/models/{m2}").status_code == 404
    assert admin.delete(f"/api/admin/vehicle-models/makes/{mk2}").json()["deleted"] is True
