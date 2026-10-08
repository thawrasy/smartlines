"""Report files: PDF, Excel, CSV, tab-separated text and JSON.

People-facing formats (PDF, Excel, CSV) carry translated headings and values in the reader's language, local
Damascus times and amounts in pounds. Integration formats (TXT, JSON) keep the stable column keys and codes so
another system can read them without a dictionary. PDF and Excel follow the brand: navy heading band, the route
line, IBM Plex Sans Arabic, and right-to-left layout in Arabic.
"""
import csv
import hashlib
import io
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from zoneinfo import ZoneInfo

from ..notify.render import messages
from .datasets import BOOL, DATE, INT, MONEY, NUM, PCT, TIME
from .engine import OutCol, Result

TZ = ZoneInfo("Asia/Damascus")
FONTS = Path(__file__).resolve().parents[2] / "assets" / "fonts"
NAVY, BLUE, LIGHT_BLUE, GREEN, INK, ROW_ALT, GRID = "#0B1F3F", "#2F7BFF", "#7FB0FF", "#12A06A", "#0F1B2D", "#F3F6FA", "#D5DCE6"
MIME = {"PDF": "application/pdf", "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "CSV": "text/csv; charset=utf-8", "TXT": "text/tab-separated-values; charset=utf-8", "JSON": "application/json"}
EXT = {"PDF": "pdf", "XLSX": "xlsx", "CSV": "csv", "TXT": "tsv", "JSON": "json"}


@dataclass
class Meta:
    code: str                 # report code, or "custom"
    title: str
    period: str
    generated_by: str
    generated_at: datetime
    locale: str


def words(locale: str) -> dict:
    w = messages(locale).get("reports") or {}
    return w if w else messages("en").get("reports", {})


def label(col: OutCol, locale: str) -> str:
    w, en = words(locale), words("en")
    base = w.get("columns", {}).get(col.label_key) or en.get("columns", {}).get(col.label_key) or col.label_key.replace("_", " ")
    if not col.agg:
        return base
    tmpl = w.get("aggregates", {}).get(col.agg) or en.get("aggregates", {}).get(col.agg, "{col}")
    return tmpl.replace("{col}", base) if col.label_key != "count" else w.get("aggregates", {}).get("count_rows", "Count")


def coded(col: OutCol, value, locale: str) -> str:
    if value is None or not col.values:
        return value
    w = messages(locale)
    group = (w.get("reports", {}).get("values", {}).get(col.values) or w.get(col.values) or {})
    return group.get(str(value), value)


def pounds(minor) -> Decimal:
    return (Decimal(int(minor)) / 100).quantize(Decimal("1")) if int(minor) % 100 == 0 else Decimal(int(minor)) / 100


def human(col: OutCol, value, locale: str):
    """A value as people read it: translated code, local time, pounds."""
    if value is None:
        return ""
    if col.type == MONEY:
        return pounds(value)
    if col.type == TIME and isinstance(value, datetime):
        return value.astimezone(TZ).replace(tzinfo=None)
    if col.type == BOOL:
        return words(locale).get("yes" if value else "no", "yes" if value else "no")
    if isinstance(value, Decimal) and col.type in (NUM, PCT):
        return float(value)
    return coded(col, value, locale)


def text(col: OutCol, value, locale: str) -> str:
    v = human(col, value, locale)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return f"{v:,}" if col.type == MONEY else str(v)
    if isinstance(v, float):
        return f"{v:g}"
    if isinstance(v, int) and col.type in (INT,):
        return f"{v:,}"
    return str(v)


def machine(col: OutCol, value):
    """A value for another system: ISO dates, UTC times, pounds as a decimal string, codes as they are."""
    if value is None:
        return None
    if col.type == MONEY:
        return str(pounds(value))
    if isinstance(value, datetime):
        return value.astimezone(ZoneInfo("UTC")).isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


# ------------------------------------------------------------------ text formats
def to_csv(res: Result, meta: Meta) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label(c, meta.locale) for c in res.columns])
    for r in res.rows:
        w.writerow([text(c, r[c.key], meta.locale).replace(",", "") if c.type == MONEY else text(c, r[c.key], meta.locale)
                    for c in res.columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")       # the BOM lets Excel open Arabic CSV correctly


def to_txt(res: Result, meta: Meta) -> bytes:
    def clean(v) -> str:
        return "" if v is None else str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")
    lines = ["\t".join(c.key for c in res.columns)]
    lines += ["\t".join(clean(machine(c, r[c.key])) for c in res.columns) for r in res.rows]
    return ("\n".join(lines) + "\n").encode("utf-8")


SOURCE_VERSION = "masslak-db-1.21.0"      # schema release the figures were computed against


def provenance(res: Result) -> dict:
    """Where a figure comes from: the moment the data reflects, the schema release, the time zone and the currency."""
    as_of = getattr(res, "data_as_of", None)
    fresh = getattr(res, "freshness", None) or {}
    return {"data_as_of": as_of.astimezone(ZoneInfo("UTC")).isoformat().replace("+00:00", "Z") if as_of else None,
            "replica_lag_seconds": fresh.get("lag_seconds"),
            "source_version": SOURCE_VERSION, "timezone": "Asia/Damascus", "currency": "SYP"}


def to_json(res: Result, meta: Meta) -> bytes:
    doc = {
        **provenance(res),
        "report": meta.code, "title": meta.title, "period": meta.period, "dataset": res.dataset,
        "generated_at": meta.generated_at.astimezone(ZoneInfo("UTC")).isoformat().replace("+00:00", "Z"),
        "columns": [{"key": c.key, "label": label(c, meta.locale), "type": c.type, **({"aggregate": c.agg} if c.agg else {})}
                    for c in res.columns],
        "rows": [{c.key: machine(c, r[c.key]) for c in res.columns} for r in res.rows],
        "totals": {k: (str(pounds(v)) if _col(res, k).type == MONEY else machine(_col(res, k), v)) for k, v in res.totals.items()},
        "truncated": res.truncated,
    }
    return json.dumps(doc, ensure_ascii=False, indent=1, default=str).encode("utf-8")


def _col(res: Result, key: str) -> OutCol:
    return next(c for c in res.columns if c.key == key)


# ------------------------------------------------------------------ Excel
def to_xlsx(res: Result, meta: Meta) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rtl = meta.locale == "ar"
    wb = Workbook()
    ws = wb.active
    ws.title = (meta.title[:28] or "report").replace("/", "-").replace(":", "")
    ws.sheet_view.rightToLeft = rtl
    wb.properties.creator = "Masslak"
    wb.properties.title = meta.title
    font = "IBM Plex Sans Arabic"
    n = max(1, len(res.columns))
    ws.cell(1, 1, meta.title).font = Font(name=font, size=15, bold=True, color=NAVY[1:])
    w = words(meta.locale)
    ws.cell(2, 1, f"{w.get('period', 'Period')}: {meta.period}").font = Font(name=font, size=10, color="5B6B80")
    ws.cell(3, 1, f"{w.get('generated', 'Generated')}: {meta.generated_at.astimezone(TZ):%Y-%m-%d %H:%M} · {meta.generated_by} · masslak.com"
            ).font = Font(name=font, size=10, color="5B6B80")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    head = 5
    thin = Side(style="thin", color=GRID[1:])
    for j, c in enumerate(res.columns, 1):
        cell = ws.cell(head, j, label(c, meta.locale))
        cell.font = Font(name=font, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY[1:])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [len(label(c, meta.locale)) + 4 for c in res.columns]
    fmt = {MONEY: "#,##0", INT: "#,##0", PCT: "0.0", NUM: "#,##0.##", DATE: "yyyy-mm-dd", TIME: "yyyy-mm-dd hh:mm"}
    alt = PatternFill("solid", fgColor=ROW_ALT[1:])
    for i, r in enumerate(res.rows):
        for j, c in enumerate(res.columns, 1):
            v = human(c, r[c.key], meta.locale)
            cell = ws.cell(head + 1 + i, j, float(v) if isinstance(v, Decimal) else v)
            cell.font = Font(name=font, size=10)
            cell.border = Border(bottom=thin)
            if c.type in fmt:
                cell.number_format = fmt[c.type]
            if i % 2:
                cell.fill = alt
            widths[j - 1] = min(48, max(widths[j - 1], len(text(c, r[c.key], meta.locale)) + 2))
    if res.totals:
        tr = head + 1 + len(res.rows)
        ws.cell(tr, 1, w.get("total", "Total")).font = Font(name=font, bold=True)
        for j, c in enumerate(res.columns, 1):
            if c.key in res.totals and j > 1:
                v = res.totals[c.key]
                cell = ws.cell(tr, j, float(pounds(v)) if c.type == MONEY else v)
                cell.font = Font(name=font, bold=True, color=NAVY[1:])
                cell.fill = PatternFill("solid", fgColor="E3EDFF")
                cell.number_format = fmt.get(c.type, "General")
    for j, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = ws.cell(head + 1, 1)
    if res.rows:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(n)}{head + len(res.rows)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ PDF
@lru_cache
def _fonts() -> tuple[str, str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("Plex", str(FONTS / "IBMPlexSansArabic-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("PlexBold", str(FONTS / "IBMPlexSansArabic-SemiBold.ttf")))
    pdfmetrics.registerFont(TTFont("PlexMono", str(FONTS / "IBMPlexMono-Medium.ttf")))
    return "Plex", "PlexBold", "PlexMono"


def _shape(s: str, rtl: bool) -> str:
    """Arabic letters joined and ordered for display; Latin-only text is left as it is."""
    if not s or not any("\u0600" <= ch <= "\u06ff" for ch in s):
        return s
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(s), base_dir="R" if rtl else "L")


def to_pdf(res: Result, meta: Meta) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    regular, bold, mono = _fonts()
    rtl = meta.locale == "ar"
    w = words(meta.locale)
    page = landscape(A4)
    align = 2 if rtl else 0
    def st(size, f=regular, col=INK, al=align):
        return ParagraphStyle("s", fontName=f, fontSize=size, leading=size * 1.45, textColor=colors.HexColor(col), alignment=al)
    cell_style = st(8)
    num_style = st(8, al=0 if rtl else 2)
    head_style = st(8, bold, "#FFFFFF", 1)
    cols = list(reversed(res.columns)) if rtl else list(res.columns)
    numeric = {MONEY, INT, NUM, PCT}

    def para(s: str, style) -> Paragraph:
        s = (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(_shape(s, rtl), style)

    data = [[para(label(c, meta.locale), head_style) for c in cols]]
    for r in res.rows:
        data.append([para(text(c, r[c.key], meta.locale), num_style if c.type in numeric else cell_style) for c in cols])
    if res.totals:
        tot = []
        for c in cols:
            if c.key in res.totals:
                v = res.totals[c.key]
                tot.append(para(f"{pounds(v):,}" if c.type == MONEY else f"{v:,}", st(8, bold, NAVY, 0 if rtl else 2)))
            else:
                tot.append(para("", cell_style))
        first = len(cols) - 1 if rtl else 0
        if not tot[first].text.strip():
            tot[first] = para(w.get("total", "Total"), st(8, bold, NAVY))
        data.append(tot)

    avail = page[0] - 28 * mm
    weights = []
    for i, c in enumerate(cols):
        longest = max([len(label(c, meta.locale))] + [len(text(c, r[c.key], meta.locale)) for r in res.rows[:300]])
        weights.append(min(40, max(6, longest)))
    total_w = sum(weights) or 1
    widths = [avail * x / total_w for x in weights]

    table = LongTable(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor(GRID)),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in range(1, len(res.rows) + 1):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(ROW_ALT)))
    if res.totals:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3EDFF")))
    table.setStyle(TableStyle(style))

    story = [Spacer(1, 4 * mm), para(meta.title, st(16, bold, NAVY)),
             para(f"{w.get('period', 'Period')}: {meta.period}   ·   {w.get('rows', 'Rows')}: {len(res.rows):,}"
                  + (f"   ·   {w.get('truncated', 'first rows only')}" if res.truncated else ""), st(9, regular, "#5B6B80")),
             Spacer(1, 4 * mm), table]
    if not res.rows:
        story.append(para(w.get("empty", "No data for this period."), st(10, regular, "#5B6B80")))

    stamp = f"{meta.generated_at.astimezone(TZ):%Y-%m-%d %H:%M}"

    def frame(canvas, doc):
        canvas.saveState()
        wd, ht = page
        canvas.setFillColor(colors.HexColor(NAVY))
        canvas.rect(0, ht - 18 * mm, wd, 18 * mm, stroke=0, fill=1)
        # the route line: light-blue origin, blue path, green destination
        x0 = (wd - 30 * mm) if rtl else 14 * mm
        canvas.setStrokeColor(colors.HexColor(BLUE))
        canvas.setLineWidth(2.6)
        canvas.setLineCap(1)
        p = canvas.beginPath()
        p.moveTo(x0, ht - 13 * mm)
        p.curveTo(x0 + 4 * mm, ht - 4 * mm, x0 + 9 * mm, ht - 4 * mm, x0 + 13 * mm, ht - 10 * mm)
        canvas.drawPath(p, stroke=1, fill=0)
        canvas.setFillColor(colors.HexColor(LIGHT_BLUE))
        canvas.circle(x0, ht - 13 * mm, 1.6 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(GREEN))
        canvas.circle(x0 + 13 * mm, ht - 10 * mm, 1.5 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont(bold, 15)
        brand = _shape(words("ar").get("brand", "Masslak"), True)
        if rtl:
            canvas.drawRightString(x0 - 4 * mm, ht - 11.5 * mm, brand)
            canvas.setFillColor(colors.HexColor("#C9D6EA"))
            canvas.setFont(regular, 10)
            canvas.drawRightString(x0 - 4 * mm - canvas.stringWidth(brand, bold, 15) - 3 * mm, ht - 11.5 * mm, "masslak")
            canvas.setFont(regular, 8)
            canvas.drawString(14 * mm, ht - 11 * mm, _shape(w.get("report_kind", "Report"), True) + "  " + meta.code)
        else:
            canvas.drawString(x0 + 18 * mm, ht - 11.5 * mm, "masslak")
            canvas.setFont(regular, 8)
            canvas.setFillColor(colors.HexColor("#C9D6EA"))
            canvas.drawRightString(wd - 14 * mm, ht - 11 * mm, f"{w.get('report_kind', 'Report')}  {meta.code}")
        # footer: who, when, page
        canvas.setFillColor(colors.HexColor("#5B6B80"))
        canvas.setFont(regular, 7.5)
        foot = _shape(f"{w.get('generated', 'Generated')}: {stamp} · {meta.generated_by}", rtl)
        page_txt = _shape(f"{w.get('page', 'Page')} {doc.page}", rtl)
        if rtl:
            canvas.drawRightString(wd - 14 * mm, 8 * mm, foot)
            canvas.drawString(14 * mm, 8 * mm, page_txt)
        else:
            canvas.drawString(14 * mm, 8 * mm, foot)
            canvas.drawRightString(wd - 14 * mm, 8 * mm, page_txt)
        canvas.setFont(mono, 7.5)
        canvas.drawCentredString(wd / 2, 8 * mm, "masslak.com")
        canvas.restoreState()

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=22 * mm, bottomMargin=14 * mm,
                            title=meta.title, author="Masslak", subject=meta.code)
    doc.build(story, onFirstPage=frame, onLaterPages=frame)
    return out.getvalue()


WRITERS = {"PDF": to_pdf, "XLSX": to_xlsx, "CSV": to_csv, "TXT": to_txt, "JSON": to_json}


def render(fmt: str, res: Result, meta: Meta) -> tuple[bytes, str]:
    """The file and its SHA-256 (kept in the export log so a handed-out file can be checked later)."""
    data = WRITERS[fmt](res, meta)
    return data, hashlib.sha256(data).hexdigest()
