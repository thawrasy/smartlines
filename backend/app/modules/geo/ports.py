"""Loading the border crossings and airports (owner's note N45, migration 1168).

The owner's approved workbook has a sheet for automatic import whose headings are English and follow the United Nations
columns (crossing_code, crossing_kind, name_ar, name_en, counterpart_name_ar, counterpart_name_en,
counterpart_country_iso2, iata_code, admin1Pcode, admin2Pcode, admin3Pcode, Location_Pcode, Latitude_y, Longitude_x,
match_quality, review_note). The sheet is found by its headings, so its name and the other sheets do not matter.

read() checks the workbook on its own; load() checks it against the geography the platform loaded and writes it in the
caller's transaction. A mistake anywhere loads nothing. A port the platform edited by hand keeps its edits; the file's
values for it come back as "kept". Every change a load makes is written to the port's history, by the file.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass
from typing import Optional

import asyncpg

from .importer import GeoFileError, _text

CODE = re.compile(r"^(LND-[A-Z]{2}-[0-9]{2}|AIR-[A-Z]{3})$")
AREA = re.compile(r"^[A-Z]{2}([0-9]{2}){1,3}$")
PLACE = re.compile(r"^[A-Z][0-9]{3,7}$")
KINDS = {"LAND_BORDER", "INTL_AIRPORT", "LOCAL_AIRPORT"}
QUALITY = {"EXACT", "TOWN", "NEAREST", "AREA"}
COLUMNS = ("crossing_code", "crossing_kind", "name_ar", "name_en", "counterpart_name_ar", "counterpart_name_en",
           "counterpart_country_iso2", "iata_code", "admin1Pcode", "admin2Pcode", "admin3Pcode", "Location_Pcode",
           "Latitude_y", "Longitude_x", "match_quality", "review_note")
# the fields a load writes, in the order of the table, and the ones a hand edit protects
FIELDS = ("name_ar", "name_en", "counterpart_name_ar", "counterpart_name_en", "counterpart_country", "iata_code",
          "area_id", "place_id", "lat", "lng", "position_source", "match_quality", "review_note")


@dataclass
class Port:
    code: str
    kind: str
    name_ar: str
    name_en: str
    counterpart_ar: Optional[str]
    counterpart_en: Optional[str]
    country: Optional[str]
    iata: Optional[str]
    sub_district: str
    place: Optional[str]
    lat: Optional[float]
    lng: Optional[float]
    quality: str
    note: Optional[str]


@dataclass
class Register:
    sha256: str
    ports: list[Port]

    def counts(self) -> dict:
        return {"ports": len(self.ports), "by_kind": {k: sum(p.kind == k for p in self.ports) for k in sorted(KINDS)}}


def _num(v) -> Optional[float]:
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _opt(v) -> Optional[str]:
    t = _text(v)
    return None if t in ("", "-", "—") else t


def read(data: bytes) -> Register:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                          # not a workbook at all
        raise GeoFileError("PORTS_FILE_FORMAT: not an Excel workbook") from e
    rows = col = None
    for ws in wb.worksheets:
        head = [_text(h) for h in next(ws.iter_rows(max_row=1, values_only=True), ())]
        if "crossing_code" in head:
            missing = [c for c in COLUMNS if c not in head]
            if missing:
                raise GeoFileError("PORTS_FILE_SHEETS: the import sheet lacks " + ", ".join(missing))
            col = {c: head.index(c) for c in COLUMNS}
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and any(v not in (None, "") for v in r)]
            break
    if rows is None:
        raise GeoFileError("PORTS_FILE_SHEETS: the workbook needs the import sheet (headed crossing_code, crossing_kind, ...)")
    problems: list[str] = []
    ports: dict[str, Port] = {}
    iatas: set[str] = set()
    for i, r in enumerate(rows, start=2):
        get = lambda c: r[col[c]] if col[c] < len(r) else None          # noqa: E731
        code, kind = _text(get("crossing_code")).upper(), _text(get("crossing_kind")).upper()
        if not CODE.match(code):
            problems.append(f"row {i}: code {code!r}")
            continue
        if code in ports:
            problems.append(f"row {i}: {code} twice")
            continue
        name_ar, name_en = _text(get("name_ar")), _text(get("name_en"))
        country, iata = (_opt(get("counterpart_country_iso2")) or "").upper() or None, (_opt(get("iata_code")) or "").upper() or None
        gov, dist, sub = (_text(get(c)).upper() for c in ("admin1Pcode", "admin2Pcode", "admin3Pcode"))
        place = (_opt(get("Location_Pcode")) or "").upper() or None
        lat, lng, quality = _num(get("Latitude_y")), _num(get("Longitude_x")), _text(get("match_quality")).upper()
        bad = [n for n, ok in (
            ("kind", kind in KINDS), ("kind and code", (kind == "LAND_BORDER") == code.startswith("LND-")),
            ("Arabic name", 2 <= len(name_ar) <= 200), ("English name", 2 <= len(name_en) <= 200),
            ("country", (country is not None and re.fullmatch("[A-Z]{2}", country) is not None) if kind == "LAND_BORDER" else country is None),
            ("country in the code", kind != "LAND_BORDER" or code[4:6] == country),
            ("IATA code", iata is None or re.fullmatch("[A-Z]{3}", iata) is not None),
            ("governorate", AREA.match(gov) is not None and len(gov) == 4),
            ("district", AREA.match(dist) is not None and len(dist) == 6),
            ("sub-district", AREA.match(sub) is not None and len(sub) == 8),
            ("place", place is None or PLACE.match(place) is not None),
            ("position", (lat is None) == (lng is None) and (lat is None or (32 <= lat <= 37.5 and 35.5 <= lng <= 42.5))),
            ("match quality", quality in QUALITY), ("place or area", place is not None or quality == "AREA")) if not ok]
        if bad:
            problems.append(f"row {i}: {code}: {', '.join(bad)}")
            continue
        if not (sub.startswith(dist) and dist.startswith(gov)):
            problems.append(f"row {i}: {code}: {sub} is not under {dist} and {gov}")
            continue
        if iata in iatas:
            problems.append(f"row {i}: {code}: IATA code {iata} twice")
            continue
        if iata:
            iatas.add(iata)
        ports[code] = Port(code, kind, name_ar, name_en, (_opt(get("counterpart_name_ar")) or None),
                           (_opt(get("counterpart_name_en")) or None), country, iata, sub, place,
                           round(lat, 6) if lat is not None else None, round(lng, 6) if lng is not None else None, quality,
                           (_opt(get("review_note")) or None))
    if problems:
        raise GeoFileError("PORTS_FILE_INVALID: " + "; ".join(problems[:12]) + (f" (and {len(problems) - 12} more)" if len(problems) > 12 else ""))
    if not ports:
        raise GeoFileError("PORTS_FILE_EMPTY: the workbook holds no crossing or airport")
    return Register(hashlib.sha256(data).hexdigest(), list(ports.values()))


async def load(conn: asyncpg.Connection, reg: Register, user_id: Optional[int], file_name: Optional[str]) -> dict:
    if await conn.fetchval("SELECT 1 FROM ref.geo_source WHERE sha256 = $1", reg.sha256):
        return {"already_loaded": True, "counts": reg.counts()}
    problems: list[str] = []
    resolved: dict[str, tuple] = {}
    for p in reg.ports:
        area = await conn.fetchval("SELECT ref.code_entity('AREA', $1)", p.sub_district)
        if area is None:
            problems.append(f"{p.code}: no sub-district {p.sub_district} (load the geography first)")
            continue
        place = None
        if p.place:
            place = await conn.fetchrow(
                """SELECT pl.id, pl.status, a.pcode AS sub_district FROM ref.place pl JOIN ref.admin_area a ON a.id = pl.area_id
                    WHERE pl.id = ref.code_entity('PLACE', $1)""", p.place)
            if place is None or place["status"] != "ACTIVE":
                problems.append(f"{p.code}: no active place {p.place}")
                continue
        if p.country and not await conn.fetchval("SELECT 1 FROM ref.country WHERE code = $1", p.country):
            problems.append(f"{p.code}: unknown country {p.country}")
            continue
        resolved[p.code] = (area, place["id"] if place else None)
    if problems:
        raise GeoFileError("PORTS_FILE_INVALID: " + "; ".join(problems[:12]) + (f" (and {len(problems) - 12} more)" if len(problems) > 12 else ""))
    source = await conn.fetchval(
        """INSERT INTO ref.geo_source (publisher, dataset, file_name, sha256, counts, imported_by, content)
           VALUES ('OWNER', 'border crossings and airports', $1, $2, $3::jsonb, $4, 'PORTS') RETURNING id""",
        (file_name or "")[:200] or None, reg.sha256, json.dumps(reg.counts()), user_id)
    await conn.execute("SELECT net.port_change_note(NULL, $1)", source)
    added = updated = 0
    kept: list[str] = []
    for p in reg.ports:
        area, place = resolved[p.code]
        row = await conn.fetchrow("SELECT id, kind, edited_by_hand FROM net.port_of_entry WHERE code = $1", p.code)
        values = (p.name_ar, p.name_en, p.counterpart_ar, p.counterpart_en, p.country, p.iata, area, place, p.lat, p.lng,
                  "UN_LOCATION" if p.lat is not None else "NONE", p.quality, (p.note or "")[:2000] or None)
        if row is None:
            await conn.execute(
                f"""INSERT INTO net.port_of_entry (code, kind, {', '.join(FIELDS)}, source_id)
                    VALUES ($1, $2, {', '.join(f'${i}' for i in range(3, 3 + len(FIELDS)))}, ${3 + len(FIELDS)})""",
                p.code, p.kind, *values, source)
            added += 1
        elif row["kind"] != p.kind:
            raise GeoFileError(f"PORTS_FILE_INVALID: {p.code} is a {row['kind']} on the platform and a {p.kind} in the file")
        elif row["edited_by_hand"]:
            kept.append(p.code)                 # the platform's edits win over a later file
        else:
            await conn.execute(
                f"""UPDATE net.port_of_entry SET {', '.join(f'{f} = ${i}' for i, f in enumerate(FIELDS, start=2))},
                        source_id = ${2 + len(FIELDS)} WHERE id = $1""", row["id"], *values, source)
            updated += 1
    await conn.execute("SELECT net.port_change_note(NULL, NULL)")
    report = reg.counts() | {"added": added, "updated": updated, "kept_edited": kept}
    await conn.execute("UPDATE ref.geo_source SET counts = $2::jsonb WHERE id = $1", source, json.dumps(report))
    return {"already_loaded": False, "counts": report}
