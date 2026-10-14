"""Loading the administrative geography from the United Nations workbook (owner's request, October 2026; migration 1124).

OCHA publishes, for each country, the common operational datasets: the administrative divisions and the populated places
with their P-codes, names in the country's language and in English, and positions. For Syria the workbook holds the
governorates (sheet admin1), districts (admin2), sub-districts (admin3), the communities and city neighbourhoods with
their positions (Location) and the areas of the cities the neighbourhoods are in (city_neighbourhoods).

read() checks the workbook and gathers it; load() writes it in the caller's transaction: what is new is added, what
changed is updated, what a newer file no longer holds is retired (never deleted: stations may stand in it), and the
cities of the network are linked to their places. The same file loaded again changes nothing. The names are written as
published; the platform's corrections of them (ref.name_correction, 1127) are applied after each load.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional

import asyncpg

AREA = re.compile(r"^[A-Z]{2}([0-9]{2}){1,3}$")
PLACE = re.compile(r"^[A-Z][0-9]{3,7}$")
# the box a country's places lie in; a position outside it is a mistake in the file
BOXES = {"SY": (32.0, 37.6, 35.4, 42.6)}
# the cities of the network as first set up, and the community each one is (the rest are linked when opened)
KNOWN_CITIES = {"DAM": "C1001", "ALP": "C1007", "HMS": "C2528", "HMA": "C2987", "LTK": "C3480", "TRT": "C5221",
                "IDL": "C3871", "DRA": "C5993", "SWD": "C6147", "DRZ": "C5086", "HSK": "C4360", "RQA": "C5710"}
SHEETS = {
    "admin1": ("admin0Pcode", "admin1Pcode", "admin1Name_en", "admin1Name_ar"),
    "admin2": ("admin1Pcode", "admin2Pcode", "admin2Name_en", "admin2Name_ar"),
    "admin3": ("admin2Pcode", "admin3Pcode", "admin3Name_en", "admin3Name_ar"),
    "Location": ("admin3Pcode", "admin4Pcode", "Location_Pcode", "LocationName_en", "LocationName_ar", "location_type",
                 "Latitude_y", "Longitude_x"),
}


class GeoFileError(ValueError):
    """The workbook is not the one expected, or holds mistakes; nothing is loaded."""


@dataclass
class Area:
    pcode: str
    level: int
    parent: Optional[str]
    name_en: str
    name_ar: str
    ref_name: Optional[str]
    valid_on: Optional[date]


@dataclass
class Place:
    pcode: str
    kind: str
    area: str
    parent: Optional[str]
    name_en: str
    name_ar: str
    lat: float
    lng: float
    city_area_code: Optional[str] = None
    city_area_en: Optional[str] = None
    city_area_ar: Optional[str] = None
    previous_pcode: Optional[str] = None
    updated_on: Optional[date] = None


@dataclass
class Workbook:
    country: str
    sha256: str
    valid_on: Optional[date]
    areas: list[Area] = field(default_factory=list)
    places: list[Place] = field(default_factory=list)

    def counts(self) -> dict:
        by = {1: 0, 2: 0, 3: 0}
        for a in self.areas:
            by[a.level] += 1
        return {"governorates": by[1], "districts": by[2], "sub_districts": by[3],
                "communities": sum(p.kind == "COMMUNITY" for p in self.places),
                "neighbourhoods": sum(p.kind == "NEIGHBOURHOOD" for p in self.places)}


BLANKS = {"", "<null>", "null", "n/a", "na", "-"}         # how the published files write an empty cell


def _text(v) -> str:
    s = " ".join(str(v).split()) if v is not None else ""
    return "" if s.lower() in BLANKS else s


def _date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _previous(v) -> Optional[str]:
    """The code a place had before, when the file names one ("new" places had none)."""
    s = _text(v)
    return s if re.fullmatch(r"[A-Z]{2}[0-9]{4,9}|[A-Z][0-9]{3,7}", s) else None


def _rows(wb, name: str, need: tuple) -> list[dict]:
    if name not in wb.sheetnames:
        raise GeoFileError(f"GEO_FILE_SHEET: the workbook has no sheet {name}")
    it = wb[name].iter_rows(values_only=True)
    head = [_text(h) for h in next(it, ())]
    missing = [c for c in need if c not in head]
    if missing:
        raise GeoFileError(f"GEO_FILE_COLUMNS: sheet {name} lacks {', '.join(missing)}")
    return [dict(zip(head, r)) for r in it if r and any(v not in (None, "") for v in r)]


def read(data: bytes) -> Workbook:
    """Reads and checks the workbook; any mistake refuses the whole file, with what is wrong and where."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                       # not a workbook at all
        raise GeoFileError("GEO_FILE_FORMAT: not an Excel workbook") from e
    a1, a2, a3, loc = (_rows(wb, n, c) for n, c in SHEETS.items())
    areas_of = {n.get("Neighbourhood_Pcode"): n for n in (_rows(wb, "city_neighbourhoods", ("Neighbourhood_Pcode",))
                                                         if "city_neighbourhoods" in wb.sheetnames else [])}
    countries = {_text(r["admin0Pcode"]) for r in a1}
    if len(countries) != 1:
        raise GeoFileError("GEO_FILE_COUNTRY: the workbook describes one country")
    country = countries.pop()
    problems: list[str] = []
    out = Workbook(country=country, sha256=hashlib.sha256(data).hexdigest(), valid_on=None)
    seen: set[str] = set()
    for level, rows, parent_col in ((1, a1, None), (2, a2, "admin1Pcode"), (3, a3, "admin2Pcode")):
        for i, r in enumerate(rows, start=2):
            code = _text(r[f"admin{level}Pcode"])
            parent = _text(r[parent_col]) if parent_col else None
            if not AREA.match(code) or len(code) != 2 + 2 * level or not code.startswith(country):
                problems.append(f"admin{level} row {i}: code {code!r}")
            elif code in seen:
                problems.append(f"admin{level} row {i}: {code} twice")
            elif parent is not None and parent not in seen:
                problems.append(f"admin{level} row {i}: {code} is under unknown {parent}")
            if not _text(r[f"admin{level}Name_en"]) or not _text(r[f"admin{level}Name_ar"]):
                problems.append(f"admin{level} row {i}: {code} has no name")
            seen.add(code)
            out.areas.append(Area(code, level, parent, _text(r[f"admin{level}Name_en"]), _text(r[f"admin{level}Name_ar"]),
                                  _text(r.get(f"admin{level}RefName") or r.get(f"admin{level}RefName_en")) or None,
                                  _date(r.get("validOn"))))
    out.valid_on = max((a.valid_on for a in out.areas if a.valid_on), default=None)
    subs = {a.pcode for a in out.areas if a.level == 3}
    box = BOXES.get(country)
    places: set[str] = set()
    for i, r in enumerate(loc, start=2):
        code, kind = _text(r["Location_Pcode"]), _text(r["location_type"]).upper()
        if kind not in ("COMMUNITY", "NEIGHBOURHOOD"):
            problems.append(f"Location row {i}: {code} is a {kind or 'blank'}")
            continue
        area, parent = _text(r["admin3Pcode"]), _text(r["admin4Pcode"])
        try:
            lat, lng = float(r["Latitude_y"]), float(r["Longitude_x"])
        except (TypeError, ValueError):
            problems.append(f"Location row {i}: {code} has no position")
            continue
        if not PLACE.match(code):
            problems.append(f"Location row {i}: code {code!r}")
        elif code in places:
            problems.append(f"Location row {i}: {code} twice")
        if area not in subs:
            problems.append(f"Location row {i}: {code} is in unknown sub-district {area}")
        if box and not (box[0] <= lat <= box[1] and box[2] <= lng <= box[3]):
            problems.append(f"Location row {i}: {code} lies outside the country ({lat}, {lng})")
        if not _text(r["LocationName_en"]) or not _text(r["LocationName_ar"]):
            problems.append(f"Location row {i}: {code} has no name")
        places.add(code)
        n = areas_of.get(code, {}) if kind == "NEIGHBOURHOOD" else {}
        out.places.append(Place(code, kind, area, parent if kind == "NEIGHBOURHOOD" else None,
                                _text(r["LocationName_en"]), _text(r["LocationName_ar"]), round(lat, 6), round(lng, 6),
                                _text(n.get("areaPcode")) or None, _text(n.get("areaName_en")) or None, _text(n.get("areaName_ar")) or None,
                                _previous(n.get("Neighbourhood_PreviuosPcode")), _date(r.get("LastUpdateDate"))))
    communities = {p.pcode for p in out.places if p.kind == "COMMUNITY"}
    for p in out.places:
        if p.kind == "NEIGHBOURHOOD" and p.parent not in communities:
            problems.append(f"Location: neighbourhood {p.pcode} is in unknown community {p.parent}")
    if problems:
        shown = "; ".join(problems[:10]) + (f"; and {len(problems) - 10} more" if len(problems) > 10 else "")
        raise GeoFileError(f"GEO_FILE_INVALID: {len(problems)} mistake(s): {shown}")
    if not out.places:
        raise GeoFileError("GEO_FILE_EMPTY: the workbook holds no places")
    return out


AREA_SQL = """
    WITH x AS (SELECT * FROM unnest($1::text[], $2::text[], $3::text[], $4::text[], $5::text[], $6::date[])
                         AS x(pcode, parent, name_en, name_ar, ref_name, valid_on))
    INSERT INTO ref.admin_area AS a (pcode, country_code, level, parent_id, name_en, name_ar, ref_name, valid_on, source_id, status)
    SELECT x.pcode, $7, $8, (SELECT id FROM ref.admin_area p WHERE p.pcode = x.parent), x.name_en, x.name_ar, x.ref_name, x.valid_on, $9, 'ACTIVE'
      FROM x
    ON CONFLICT (pcode) DO UPDATE SET parent_id = excluded.parent_id, name_en = excluded.name_en, name_ar = excluded.name_ar,
        ref_name = excluded.ref_name, valid_on = excluded.valid_on, source_id = excluded.source_id, status = 'ACTIVE'
     WHERE (a.parent_id, a.published_en, a.published_ar, a.ref_name, a.valid_on, a.status)
           IS DISTINCT FROM (excluded.parent_id, excluded.name_en, excluded.name_ar, excluded.ref_name, excluded.valid_on, 'ACTIVE')
    RETURNING (xmax = 0) AS added"""

PLACE_SQL = """
    WITH x AS (SELECT * FROM unnest($1::text[], $2::text[], $3::text[], $4::text[], $5::text[], $6::text[], $7::numeric[],
                                    $8::numeric[], $9::text[], $10::text[], $11::text[], $12::text[], $13::date[])
                         AS x(pcode, kind, area, parent, name_en, name_ar, lat, lng, ca_code, ca_en, ca_ar, previous, updated_on))
    INSERT INTO ref.place AS p (pcode, kind, area_id, parent_place_id, name_en, name_ar, lat, lng, city_area_code, city_area_en,
                                city_area_ar, previous_pcode, updated_on, source_id, status)
    SELECT x.pcode, x.kind, (SELECT id FROM ref.admin_area a WHERE a.pcode = x.area),
           (SELECT id FROM ref.place c WHERE c.pcode = x.parent), x.name_en, x.name_ar, x.lat, x.lng, x.ca_code, x.ca_en, x.ca_ar,
           x.previous, x.updated_on, $14, 'ACTIVE'
      FROM x
    ON CONFLICT (pcode) DO UPDATE SET kind = excluded.kind, area_id = excluded.area_id, parent_place_id = excluded.parent_place_id,
        name_en = excluded.name_en, name_ar = excluded.name_ar, lat = excluded.lat, lng = excluded.lng,
        city_area_code = excluded.city_area_code, city_area_en = excluded.city_area_en, city_area_ar = excluded.city_area_ar,
        previous_pcode = excluded.previous_pcode, updated_on = excluded.updated_on, source_id = excluded.source_id, status = 'ACTIVE'
     WHERE (p.kind, p.area_id, p.parent_place_id, p.published_en, p.published_ar, p.lat, p.lng, p.city_area_code, p.status)
           IS DISTINCT FROM (excluded.kind, excluded.area_id, excluded.parent_place_id, excluded.name_en, excluded.name_ar,
                             excluded.lat, excluded.lng, excluded.city_area_code, 'ACTIVE')
    RETURNING (xmax = 0) AS added"""


async def load(conn: asyncpg.Connection, wb: Workbook, user_id: Optional[int], file_name: Optional[str]) -> dict:
    """Writes the workbook (the platform's context, in the caller's transaction) and reports what changed."""
    if await conn.fetchval("SELECT 1 FROM ref.geo_source WHERE sha256 = $1", wb.sha256):
        return {"already_loaded": True, "counts": wb.counts()}
    if not await conn.fetchval("SELECT 1 FROM ref.country WHERE code = $1", wb.country):
        raise GeoFileError(f"GEO_FILE_COUNTRY: country {wb.country} is not known to the platform")
    source = await conn.fetchval(
        """INSERT INTO ref.geo_source (publisher, dataset, file_name, sha256, valid_on, counts, imported_by)
           VALUES ('OCHA', $1, $2, $3, $4, $5::jsonb, $6) RETURNING id""",
        f"{wb.country} administrative divisions and populated places (P-codes)", (file_name or "")[:200] or None,
        wb.sha256, wb.valid_on, json.dumps(wb.counts()), user_id)
    added = updated = 0
    for level in (1, 2, 3):
        rows = [a for a in wb.areas if a.level == level]
        changed = await conn.fetch(AREA_SQL, [a.pcode for a in rows], [a.parent for a in rows], [a.name_en for a in rows],
                                   [a.name_ar for a in rows], [a.ref_name for a in rows], [a.valid_on for a in rows],
                                   wb.country, level, source)
        added += sum(r["added"] for r in changed)
        updated += sum(not r["added"] for r in changed)
    for kind in ("COMMUNITY", "NEIGHBOURHOOD"):                   # the cities first: neighbourhoods name them
        rows = [p for p in wb.places if p.kind == kind]
        changed = await conn.fetch(PLACE_SQL, [p.pcode for p in rows], [p.kind for p in rows], [p.area for p in rows],
                                   [p.parent for p in rows], [p.name_en for p in rows], [p.name_ar for p in rows],
                                   [p.lat for p in rows], [p.lng for p in rows], [p.city_area_code for p in rows],
                                   [p.city_area_en for p in rows], [p.city_area_ar for p in rows], [p.previous_pcode for p in rows],
                                   [p.updated_on for p in rows], source)
        added += sum(r["added"] for r in changed)
        updated += sum(not r["added"] for r in changed)
    # what this country's published file no longer holds is retired; places surveyed in the field stay
    retired = await conn.fetchval(
        """WITH r AS (UPDATE ref.place p SET status = 'RETIRED'
                       WHERE p.status = 'ACTIVE' AND NOT (p.pcode = ANY($1::text[]))
                         AND p.area_id IN (SELECT id FROM ref.admin_area WHERE country_code = $2)
                         AND (p.source_id IS NULL OR p.source_id IN (SELECT id FROM ref.geo_source WHERE publisher = 'OCHA'))
                       RETURNING 1)
           SELECT count(*) FROM r""", [p.pcode for p in wb.places], wb.country)
    retired += await conn.fetchval(
        """WITH r AS (UPDATE ref.admin_area SET status = 'RETIRED'
                       WHERE status = 'ACTIVE' AND country_code = $2 AND NOT (pcode = ANY($1::text[])) RETURNING 1)
           SELECT count(*) FROM r""", [a.pcode for a in wb.areas], wb.country)
    linked = await conn.fetchval(
        """WITH k AS (SELECT * FROM unnest($1::text[], $2::text[]) AS k(code, pcode)),
                l AS (UPDATE ref.city c SET place_id = p.id FROM k JOIN ref.place p ON p.pcode = k.pcode
                       WHERE c.code = k.code AND c.place_id IS NULL
                         AND NOT EXISTS (SELECT 1 FROM ref.city o WHERE o.place_id = p.id) RETURNING 1)
           SELECT count(*) FROM l""", list(KNOWN_CITIES), list(KNOWN_CITIES.values()))
    # the platform's corrections of the published names hold for the newer file too (1127)
    await conn.fetchval("SELECT ref.apply_name_corrections()")
    report = wb.counts() | {"added": added, "updated": updated, "retired": retired, "cities_linked": linked}
    await conn.execute("UPDATE ref.geo_source SET counts = $2::jsonb WHERE id = $1", source, json.dumps(report))
    return {"already_loaded": False, "counts": report}
