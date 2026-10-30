"""Builds docs/integration/saudi-iban/: the Saudi IBAN bank-code workbook with a live ISO 13616 mod-97 validator, and
saudi_banks.json, the same bank table for the reference validators beside it.

Every check digit, bank name and verdict in the workbook is an Excel formula, so the sheet recalculates when an IBAN or an
account number is typed in. Run `python scripts/gen_saudi_iban_workbook.py`, then recalculate the file in LibreOffice or Excel.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PACKAGE = Path(__file__).resolve().parent.parent / "docs" / "integration" / "saudi-iban"
OUT = PACKAGE / "saudi_iban_bank_codes.xlsx"

# (code, Arabic name, English name, status, note). Source: SAMA-assigned IBAN bank identifiers as published in
# github.com/EngMEita/saudi-banks and cross-checked with cubbbix.com and validatefin.com (October 2026).
BANKS = [
    ("05", "مصرف الإنماء", "Alinma Bank", "نشط", ""),
    ("10", "البنك الأهلي السعودي", "Saudi National Bank (SNB)", "نشط", "يشمل حسابات سامبا المنقولة بعد الاندماج 2021"),
    ("15", "بنك البلاد", "Bank Albilad", "نشط", ""),
    ("20", "بنك الرياض", "Riyad Bank", "نشط", ""),
    ("30", "البنك العربي الوطني", "Arab National Bank", "نشط", ""),
    ("36", "بنك دال ثلاثمائة وستون (D360)", "D360 Bank", "بنك رقمي",
     "مؤكد بآيبان فعلي لحساب في البنك قدّمه المستخدم (SA83 3603 …) يجتاز MOD-97، ويتفق مع الملف الاسترشادي"),
    ("40", "مجموعة سامبا المالية", "Samba Financial Group", "مدمج (تاريخي)", "اندمج في البنك الأهلي السعودي 2021؛ قد تبقى آيبانات قديمة متداولة"),
    ("45", "البنك السعودي الأول (ساب)", "Saudi Awwal Bank (SAB, formerly SABB)", "نشط", ""),
    ("50", "البنك الأول", "Alawwal Bank", "مدمج (تاريخي)", "اندمج في ساب 2019؛ قد تبقى آيبانات قديمة متداولة"),
    ("55", "البنك السعودي الفرنسي", "Banque Saudi Fransi", "نشط", ""),
    ("60", "بنك الجزيرة", "Bank AlJazira", "نشط", ""),
    ("65", "البنك السعودي للاستثمار", "The Saudi Investment Bank", "نشط", ""),
    ("71", "بنك البحرين الوطني (فرع السعودية)", "National Bank of Bahrain", "فرع بنك أجنبي", ""),
    ("73", "بنك أبوظبي الأول (فرع السعودية)", "First Abu Dhabi Bank (FAB)", "فرع بنك أجنبي",
     "مؤكد بآيبان فعلي لحساب في البنك قدّمه المستخدم (SA63 7300 …) يجتاز MOD-97، ويتفق مع الملف الاسترشادي"),
    ("75", "بنك الكويت الوطني (فرع السعودية)", "National Bank of Kuwait", "فرع بنك أجنبي", ""),
    ("76", "بنك مسقط (فرع السعودية)", "Bank Muscat", "فرع بنك أجنبي", ""),
    ("78", "بنك إس تي سي (STC Bank)", "STC Bank", "بنك رقمي", "مصدر واحد: آيبان منشور SA47 7800 0000 0011 8115 9355 يجتاز MOD-97 "
     "— يُستحسن التأكد من ساما"),
    ("80", "مصرف الراجحي", "Al Rajhi Bank", "نشط", ""),
    ("81", "دويتشه بنك (فرع السعودية)", "Deutsche Bank", "فرع بنك أجنبي", ""),
    ("82", "البنك الوطني الباكستاني (فرع السعودية)", "National Bank of Pakistan", "فرع بنك أجنبي", ""),
    ("83", "بنك الدولة الهندي (فرع السعودية)", "State Bank of India", "فرع بنك أجنبي", "مصدر واحد فقط — يُستحسن التأكد من ساما"),
    ("84", "بنك زراعات التركي (فرع السعودية)", "T.C. Ziraat Bankası", "فرع بنك أجنبي", ""),
    ("85", "بي إن بي باريبا (فرع السعودية)", "BNP Paribas", "فرع بنك أجنبي", ""),
    ("86", "جي بي مورغان تشيس (فرع السعودية)", "JPMorgan Chase Bank", "فرع بنك أجنبي", ""),
    ("90", "بنك الخليج الدولي", "Gulf International Bank (GIB)", "فرع بنك أجنبي", ""),
    ("95", "بنك الإمارات دبي الوطني (فرع السعودية)", "Emirates NBD", "فرع بنك أجنبي", ""),
]

# Head-office SWIFT/BIC per IBAN bank code, with how well it is attested in public BIC directories (October 2026).
CONFIRMED, ONE_SOURCE, UNCONFIRMED = "مؤكد", "مصدر واحد", "غير مؤكد"
SWIFT = {
    "05": ("INMASARI", CONFIRMED), "10": ("NCBKSAJE", CONFIRMED), "15": ("ALBISARI", CONFIRMED),
    "20": ("RIBLSARI", CONFIRMED), "30": ("ARNBSARI", CONFIRMED), "36": ("DBAKSARI", CONFIRMED),
    "40": ("SAMBSARI", UNCONFIRMED + " (بنك مدمج)"), "45": ("SABBSARI", CONFIRMED),
    "50": ("AAALSARI", ONE_SOURCE + " (بنك مدمج)"), "55": ("BSFRSARI", CONFIRMED), "60": ("BJAZSAJE", CONFIRMED),
    "65": ("SIBCSARI", CONFIRMED), "71": ("NBOBSARI", ONE_SOURCE), "73": ("FABMSARI", ONE_SOURCE + " + الملف الاسترشادي"),
    "75": ("NBOKSAJE", ONE_SOURCE + " + الملف الاسترشادي"), "76": ("BMUSSARI", ONE_SOURCE), "78": ("STCJSARI", CONFIRMED),
    "80": ("RJHISARI", CONFIRMED), "81": ("DEUTSARI", ONE_SOURCE), "82": ("NBPASARI", CONFIRMED + " (ورد NBOPSARI في جدول واحد)"),
    "83": ("SBINSAJE", ONE_SOURCE), "84": ("TCZBSAJE", ONE_SOURCE), "85": ("BNPASARI", UNCONFIRMED + " (الملف الاسترشادي فقط)"),
    "86": ("CHASSARI", UNCONFIRMED + " (ورد CHASSARISAU لشركة JPMorgan السعودية)"),
    "90": ("GULFSARI", UNCONFIRMED + " (الملف الاسترشادي فقط)"), "95": ("EBILSARI", ONE_SOURCE + " + الملف الاسترشادي"),
}

# The SWIFT IBAN Registry's example account for Saudi Arabia (SA03 8000 0000 6080 1016 7519): reused for every bank so
# the table shows that the check digits change with the bank code alone.
EXAMPLE_ACCOUNT = "000000608010167519"

FONT = "Arial"
NAVY = "1F3864"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", start_color=NAVY)
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
FILL_ENGINE = PatternFill("solid", start_color="F2F2F2")
GREEN = PatternFill("solid", start_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(horizontal="right", vertical="center", wrap_text=True)

BANKS_SHEET = "رموز البنوك"
BANK_RANGE = f"'{BANKS_SHEET}'!$B$5:$B${4 + len(BANKS)}"
NAME_RANGE = f"'{BANKS_SHEET}'!$C$5:$C${4 + len(BANKS)}"


def mod97_chain(src: str, n_chunks: int) -> str:
    """A formula for MOD(src, 97) on a digit string longer than Excel's 15-digit precision: 9 digits, then 7 at a time."""
    expr = f"MOD(VALUE(MID({src},1,9)),97)"
    for i in range(n_chunks):
        expr = f"MOD(VALUE({expr}&MID({src},{10 + 7 * i},7)),97)"
    return expr


def header(ws, row, titles, widths=None):
    for col, title in enumerate(titles, start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BOX
    for col, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 32


def style_row(ws, row, ncols, align=CENTER):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.border, c.alignment = BOX, align
        if c.font == Font():
            c.font = F_BASE


def build_structure(wb):
    ws = wb.active
    ws.title = "هيكل الآيبان"
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "هيكل رقم الآيبان السعودي (ISO 13616) — 24 خانة"
    ws["A1"].font = F_TITLE
    ws["A2"] = "SA + خانتا تدقيق + رمز البنك (خانتان) + رقم الحساب (18 خانة). المرجع: سجل SWIFT للآيبان، صيغة SA2!n2!n18!c"
    ws["A2"].font = F_NOTE
    header(ws, 4, ["الجزء", "عدد الخانات", "المواقع", "مثال", "الوصف"], [18, 12, 12, 22, 70])
    rows = [
        ("رمز الدولة", 2, "1–2", "SA", "يرمز للمملكة العربية السعودية (ISO 3166). ثابت لكل الحسابات السعودية."),
        ("حقل التدقيق", 2, "3–4", "03", "رقمان يُحسبان بخوارزمية MOD-97 من رمز البنك ورقم الحساب معاً، فهما يختلفان من حساب لآخر "
         "حتى داخل البنك نفسه، ولا يوجد حقل تدقيق ثابت للبنك. وظيفتهما كشف أخطاء الإدخال (أي خطأ في خانة واحدة أو تبديل خانتين متجاورتين)."),
        ("رمز البنك", 2, "5–6", "80", "رقم ثنائي تصدره ساما لكل بنك: 80 مصرف الراجحي، 10 البنك الأهلي السعودي، 20 بنك الرياض، "
         "55 البنك السعودي الفرنسي. الجدول الكامل في ورقة «رموز البنوك»."),
        ("رقم الحساب", 18, "7–24", "000000608010167519", "رقم الحساب لدى البنك، يُكمَّل بأصفار على اليسار حتى 18 خانة (يسمح المعيار بحروف وأرقام)."),
        ("الآيبان كاملاً", 24, "1–24", "SA03 8000 0000 6080 1016 7519", "يُكتب إلكترونياً بلا مسافات، ويُطبع في مجموعات من 4 خانات."),
    ]
    for i, r in enumerate(rows, start=5):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
        style_row(ws, i, 5)
        ws.cell(row=i, column=5).alignment = WRAP
        ws.row_dimensions[i].height = 48
    ws["A11"] = "تصحيح للأمثلة الواردة في الطلب: الرمز 20 هو بنك الرياض (وليس الراجحي)، و55 هو البنك السعودي الفرنسي (وليس الأهلي). " \
                "الراجحي = 80، الأهلي = 10."
    ws["A11"].font = Font(name=FONT, size=10, bold=True, color="C00000")
    ws.merge_cells("A11:E11")
    ws.row_dimensions[11].height = 30
    ws["A11"].alignment = WRAP


def build_banks(wb):
    ws = wb.create_sheet(BANKS_SHEET)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "رموز البنوك العاملة في المملكة العربية السعودية في رقم الآيبان"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("حقل التدقيق ليس صفة للبنك بل للحساب: العمود J مدخل (رقم حساب من 18 خانة)، والعمودان K و L يحسبان حقل التدقيق والآيبان آلياً. "
                "القيمة الافتراضية في J هي حساب المثال في سجل SWIFT، مكررة لكل بنك لإظهار تغيّر خانتي التدقيق بتغيّر رمز البنك.")
    ws["A2"].font = F_NOTE
    ws.merge_cells("A2:M2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 30
    header(ws, 4, ["#", "رمز البنك", "اسم البنك", "Bank name", "رمز SWIFT / BIC", "توثيق SWIFT", "الحالة", "ملاحظة",
                   "الآيبان يبدأ بـ", "رقم الحساب (18 خانة) — مدخل", "حقل التدقيق (محسوب)", "الآيبان الكامل (محسوب)",
                   "التحقق MOD-97", "سلسلة الحساب (مساعد)"],
           [5, 10, 34, 34, 14, 22, 15, 40, 14, 24, 14, 34, 14, 30])
    for i, (code, ar, en, status, note) in enumerate(BANKS, start=5):
        swift, level = SWIFT[code]
        ws.cell(row=i, column=1, value=i - 4)
        ws.cell(row=i, column=2, value=code).number_format = "@"
        ws.cell(row=i, column=3, value=ar)
        ws.cell(row=i, column=4, value=en)
        ws.cell(row=i, column=5, value=swift)
        ws.cell(row=i, column=6, value=level)
        ws.cell(row=i, column=7, value=status)
        ws.cell(row=i, column=8, value=note)
        ws.cell(row=i, column=9, value=f'="SA"&"xx"&B{i}')
        acct = ws.cell(row=i, column=10, value=EXAMPLE_ACCOUNT)
        acct.number_format, acct.font, acct.fill = "@", F_INPUT, FILL_INPUT
        # Helper N: bank code + account + "SA00" as digits (S=28, A=10). Saudi accounts are digits in practice.
        ws.cell(row=i, column=14, value=f'=B{i}&J{i}&"281000"')
        ws.cell(row=i, column=11, value=f'=IF(LEN(J{i})<>18,"طول خاطئ",TEXT(98-{mod97_chain(f"N{i}", 3)},"00"))')
        ws.cell(row=i, column=12, value=f'=IF(LEN(J{i})<>18,"","SA"&K{i}&" "&MID(B{i}&J{i},1,4)&" "&MID(B{i}&J{i},5,4)&" "'
                                        f'&MID(B{i}&J{i},9,4)&" "&MID(B{i}&J{i},13,4)&" "&MID(B{i}&J{i},17,4))')
        ws.cell(row=i, column=13, value=f'=IF(LEN(J{i})<>18,"",IF({mod97_chain(f"(B{i}&J{i}&\"2810\"&K{i})", 3)}=1,"صحيح ✓","خطأ"))')
        style_row(ws, i, 14)
        for col in (3, 4, 8):
            ws.cell(row=i, column=col).alignment = WRAP
        ws.cell(row=i, column=14).fill = FILL_ENGINE
        if status.startswith("مدمج"):
            for col in range(1, 9):
                ws.cell(row=i, column=col).font = Font(name=FONT, size=10, color="808080")
    ws.conditional_formatting.add(f"F5:F{4 + len(BANKS)}", FormulaRule(formula=['LEFT($F5,4)="مؤكد"'], fill=GREEN))
    ws.conditional_formatting.add(f"F5:F{4 + len(BANKS)}", FormulaRule(formula=['LEFT($F5,3)="غير"'], fill=FILL_INPUT))
    last = 4 + len(BANKS)
    ws.cell(row=last + 2, column=1, value="المصادر: معرّفات البنوك التي تصدرها ساما كما ينشرها github.com/EngMEita/saudi-banks، مع مراجعة "
                                          "cubbbix.com و validatefin.com (أكتوبر 2026). الرمزان 78 (STC Bank) و83 وردا في مصدر واحد. "
                                          "الرمزان 36 (D360) و73 (بنك أبوظبي الأول) مؤكدان بآيبانات فعلية قدّمها المستخدم. "
                                          "Vision Bank لم يُدرج لأن رمزه لم يُتحقق منه؛ يُعرف من الخانتين 5–6 في أي آيبان صادر منه ويُضاف في صف جديد. "
                                          "رموز SWIFT: ثماني خانات للمركز الرئيسي (تُضاف XXX لصيغة 11 خانة)، من أدلة SWIFT العامة (Wise، XE، theswiftcodes، "
                                          "ksaexpats، alaan) — «مؤكد» = أكثر من مصدر، «مصدر واحد» و«غير مؤكد» يُتحقق منهما مع البنك قبل التحويل.")
    ws.cell(row=last + 2, column=1).font = F_NOTE
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=13)
    ws.cell(row=last + 2, column=1).alignment = WRAP
    ws.row_dimensions[last + 2].height = 60
    ws.cell(row=last + 3, column=1, value="مفتاح الألوان: الخلايا الصفراء بخط أزرق مدخلات يمكن تعديلها؛ الباقي معادلات.").font = F_NOTE
    ws.freeze_panes = "C5"
    ws.cell(row=4, column=11).comment = Comment("حقل التدقيق = 98 − MOD(رمز البنك & رقم الحساب & \"SA00\" بعد تحويل الحروف، 97)", "IBAN")


def build_validator(wb, rows=40):
    ws = wb.create_sheet("التحقق من الآيبان")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "التحقق من صحة الآيبان السعودي — خوارزمية ISO 7064 MOD 97-10"
    ws["A1"].font = F_TITLE
    ws["A2"] = "اكتب أو الصق الآيبان في العمود A (بمسافات أو بدونها). النتيجة والسبب وحقل التدقيق الصحيح تظهر آلياً. الأعمدة من L فما بعد محرك الحساب."
    ws["A2"].font = F_NOTE
    ws.merge_cells("A2:J2")
    titles = ["الآيبان (مدخل)", "بعد التنظيف", "الطول", "الدولة", "حقل التدقيق المُدخل", "رمز البنك", "اسم البنك",
              "باقي القسمة على 97", "حقل التدقيق الصحيح لهذا الحساب", "النتيجة", "السبب"]
    header(ws, 4, titles, [32, 28, 8, 8, 12, 10, 32, 12, 16, 12, 44])

    first_char_col = 14                      # N: first of 34 per-character columns
    num_col = first_char_col + 34            # the whole IBAN as a digit string
    chunk0 = num_col + 1                     # running remainders, 9 digits then 7 at a time
    n_chunks = 10                            # 9 + 9*7 = 72 >= 68 digits (34 letters at most)
    eng = ["المُعاد ترتيبه", "صالح الرموز"] + [f"ح{k}" for k in range(1, 35)] + ["السلسلة الرقمية"] + \
          [f"باقي {k}" for k in range(1, n_chunks + 1)]
    for off, t in enumerate(eng):
        c = ws.cell(row=4, column=12 + off, value=t)
        c.font, c.fill, c.alignment, c.border = F_BOLD, FILL_ENGINE, CENTER, BOX
        ws.column_dimensions[get_column_letter(12 + off)].width = 6 if 0 < off < 36 else 16
    ws.column_dimensions["L"].width = 30
    ws.column_dimensions[get_column_letter(num_col)].width = 40

    L = lambda col, r: f"{get_column_letter(col)}{r}"
    examples = ["SA03 8000 0000 6080 1016 7519",       # the SWIFT registry example: valid
                "SA04 8000 0000 6080 1016 7519",       # one wrong check digit
                "SA03 8000 0000 6080 1016 7591",       # two digits swapped
                "SA44 2000 0001 2345 6789 1234",       # Riyad Bank, valid
                "SA03 8000 0000 6080 1016 751",        # one digit short
                "SA33 9900 0000 6080 1016 7519",       # valid mod-97 but no such bank code
                "AE07 0331 2345 6789 0123 456",        # not a Saudi IBAN
                "SA47 7800 0000 0011 8115 9355",       # STC Bank, published example: valid
                "SA43 7300 0000 6080 1016 7519",       # First Abu Dhabi Bank, registry account: valid
                "sa05 3600 0000 6080 1016 7519"]       # D360 Bank in lower case, registry account: valid
    for r in range(5, 5 + rows):
        a = ws.cell(row=r, column=1, value=examples[r - 5] if r - 5 < len(examples) else None)
        a.font, a.fill, a.number_format = F_INPUT, FILL_INPUT, "@"
        ws[f"B{r}"] = f'=UPPER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM(A{r})," ",""),"-",""),CHAR(160),""))'
        ws[f"C{r}"] = f'=IF(B{r}="","",LEN(B{r}))'
        ws[f"D{r}"] = f'=IF(B{r}="","",LEFT(B{r},2))'
        ws[f"E{r}"] = f'=IF(B{r}="","",MID(B{r},3,2))'
        ws[f"F{r}"] = f'=IF(B{r}="","",MID(B{r},5,2))'
        ws[f"G{r}"] = f'=IF(B{r}="","",IFERROR(INDEX({NAME_RANGE},MATCH(F{r},{BANK_RANGE},0)),"رمز غير معروف"))'
        # Engine: move the first four characters to the end, turn A..Z into 10..35, take the remainder in chunks.
        ws[f"L{r}"] = f'=IF(B{r}="","",MID(B{r},5,40)&LEFT(B{r},4))'
        for k in range(1, 35):
            ch = f"MID($L{r},{k},1)"
            ws.cell(row=r, column=first_char_col + k - 1,
                    value=f'=IF(LEN($L{r})<{k},"",IF(ISNUMBER(FIND({ch},"0123456789")),{ch},'
                          f'IF(ISNUMBER(FIND({ch},"ABCDEFGHIJKLMNOPQRSTUVWXYZ")),CODE({ch})-55,"#")))')
        ws[f"M{r}"] = f'=IF(L{r}="","",NOT(ISNUMBER(FIND("#",{L(num_col, r)}))))'
        ws.cell(row=r, column=num_col,
                value=f"=_xlfn.CONCAT({L(first_char_col, r)}:{L(first_char_col + 33, r)})")
        s = L(num_col, r)
        ws.cell(row=r, column=chunk0, value=f'=IF(OR({s}="",M{r}=FALSE),"",MOD(VALUE(MID({s},1,9)),97))')
        for i in range(1, n_chunks):
            prev = L(chunk0 + i - 1, r)
            ws.cell(row=r, column=chunk0 + i, value=f'=IF({prev}="","",MOD(VALUE({prev}&MID({s},{10 + 7 * (i - 1)},7)),97))')
        rem = L(chunk0 + n_chunks - 1, r)
        ws[f"H{r}"] = f"={rem}"
        digits_ok = f'AND(ISNUMBER(FIND(MID(B{r},3,1),"0123456789")),ISNUMBER(FIND(MID(B{r},4,1),"0123456789")))'
        # The check digits that would make this account valid: 98 - (N - cd) mod 97, since N = 100*X + cd.
        ws[f"I{r}"] = f'=IF(OR(H{r}="",NOT({digits_ok})),"",TEXT(98-MOD(H{r}-VALUE(E{r}),97),"00"))'
        ws[f"K{r}"] = (f'=IF(B{r}="","",'
                       f'IF(M{r}=FALSE,"يحتوي رموزاً غير مسموحة (المسموح A–Z و 0–9 فقط)",'
                       f'IF(D{r}<>"SA","رمز الدولة "&D{r}&" — هذا المدقق للآيبان السعودي SA",'
                       f'IF(C{r}<>24,"الطول "&C{r}&" خانة والمطلوب 24",'
                       f'IF(NOT({digits_ok}),"خانتا التدقيق يجب أن تكونا رقمين",'
                       f'IF(ISERROR(VALUE(MID(B{r},5,20))),"رمز البنك ورقم الحساب يجب أن يكونا أرقاماً",'
                       f'IF(H{r}<>1,"فشل اختبار MOD-97 (الباقي "&H{r}&" وليس 1) — خطأ في كتابة الرقم؛ خانتا التدقيق الصحيحتان لهذا الحساب "&I{r},'
                       f'IF(G{r}="رمز غير معروف","اجتاز MOD-97 لكن رمز البنك "&F{r}&" غير مسجل في جدول البنوك",'
                       f'"آيبان صحيح — "&G{r}))))))))')
        ws[f"J{r}"] = f'=IF(B{r}="","",IF(LEFT(K{r},10)="آيبان صحيح","صحيح ✓","غير صحيح ✗"))'
        style_row(ws, r, 11)
        ws[f"A{r}"].font = F_INPUT
        ws[f"K{r}"].alignment = WRAP
        for col in range(12, chunk0 + n_chunks):
            ws.cell(row=r, column=col).font = Font(name=FONT, size=8, color="808080")
            ws.cell(row=r, column=col).fill = FILL_ENGINE
    end = 4 + rows
    ws.conditional_formatting.add(f"J5:J{end}", FormulaRule(formula=['LEFT($J5,4)="صحيح"'], fill=GREEN))
    ws.conditional_formatting.add(f"J5:J{end}", FormulaRule(formula=['LEFT($J5,3)="غير"'], fill=RED))
    for col in range(13, chunk0 + n_chunks):        # keep L visible, fold the per-character engine away
        ws.column_dimensions[get_column_letter(col)].outlineLevel = 1
        ws.column_dimensions[get_column_letter(col)].hidden = True
    ws.freeze_panes = "B5"
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="42", allow_blank=True,
                        promptTitle="الآيبان", prompt="مثال: SA03 8000 0000 6080 1016 7519")
    ws.add_data_validation(dv)
    dv.add(f"A5:A{end}")


def build_algorithm(wb):
    ws = wb.create_sheet("شرح الخوارزمية")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "خوارزمية MOD-97 خطوة بخطوة (ISO 13616 / ISO 7064)"
    ws["A1"].font = F_TITLE
    ws["A2"] = "غيّر الآيبان في الخلية C4 لترى كل خطوة تُحسب عليه. (مثال حي يعمل للآيبان السعودي المكوّن من أرقام)"
    ws["A2"].font = F_NOTE
    header(ws, 3, ["الخطوة", "الشرح", "النتيجة"], [8, 62, 44])
    ws.row_dimensions[3].height = 24
    steps = [
        ("0", "الآيبان كما أُدخل", "SA03 8000 0000 6080 1016 7519"),
        ("1", "إزالة المسافات وتحويل الحروف إلى كبيرة، والتأكد أن الطول 24 للسعودية", '=UPPER(SUBSTITUTE(C4," ",""))'),
        ("1أ", "الطول", "=LEN(C5)"),
        ("2", "نقل أول 4 خانات (رمز الدولة + حقل التدقيق) إلى نهاية الرقم", "=MID(C5,5,40)&LEFT(C5,4)"),
        ("3", "استبدال كل حرف برقمه: A=10، B=11 … S=28 … Z=35  (SA ← 2810)",
         '=MID(C7,1,20)&(CODE(MID(C7,21,1))-55)&(CODE(MID(C7,22,1))-55)&MID(C7,23,2)'),
        ("4", "باقي أول 9 خانات على 97", "=MOD(VALUE(MID(C8,1,9)),97)"),
        ("5", "الباقي السابق + الخانات 7 التالية، ثم الباقي على 97", "=MOD(VALUE(C9&MID(C8,10,7)),97)"),
        ("6", "الباقي السابق + الخانات 7 التالية، ثم الباقي على 97", "=MOD(VALUE(C10&MID(C8,17,7)),97)"),
        ("7", "الباقي السابق + الخانات المتبقية، ثم الباقي على 97 = باقي الرقم كاملاً", "=MOD(VALUE(C11&MID(C8,24,7)),97)"),
        ("8", "الحكم: الآيبان صحيح إذا وفقط إذا كان الباقي = 1", '=IF(AND(C6=24,C12=1),"صحيح ✓","غير صحيح ✗")'),
        ("", "", ""),
        ("حساب", "كيف تُولَّد خانتا التدقيق لحساب جديد:", ""),
        ("أ", "ضع 00 مكان خانتي التدقيق: رمز البنك & رقم الحساب & SA00", '=MID(C5,5,20)&"281000"'),
        ("ب", "احسب الباقي على 97 بالطريقة نفسها (على دفعات)", "=" + mod97_chain("C16", 3)),
        ("ج", "خانتا التدقيق = 98 − الباقي (بخانتين، مع صفر بادئ)", '=TEXT(98-C17,"00")'),
        ("د", "الآيبان الناتج", '="SA"&C18&MID(C5,5,20)'),
    ]
    for i, (n, text, val) in enumerate(steps, start=4):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=text)
        c = ws.cell(row=i, column=3, value=val)
        if n:
            style_row(ws, i, 3)
            ws.cell(row=i, column=2).alignment = WRAP
        if i in (7, 8, 16):
            c.number_format = "@"
    ws["C4"].font, ws["C4"].fill = F_INPUT, FILL_INPUT
    ws["B15"].font = F_BOLD
    ws.conditional_formatting.add("C13", FormulaRule(formula=['LEFT($C$13,4)="صحيح"'], fill=GREEN))
    ws.conditional_formatting.add("C13", FormulaRule(formula=['LEFT($C$13,3)="غير"'], fill=RED))
    notes = [
        "لماذا على دفعات؟ الرقم الناتج 26 خانة، ودقة Excel 15 خانة فقط، فيُقسم إلى أجزاء لا تتجاوز 9 خانات مع حمل الباقي — النتيجة مطابقة رياضياً.",
        "ما الذي تكشفه الخوارزمية؟ كل خطأ في خانة واحدة، وكل تبديل لخانتين متجاورتين، ونحو 99% من الأخطاء الأخرى.",
        "ما الذي لا تكشفه؟ صحة الآيبان رياضياً لا تعني أن الحساب موجود أو مفتوح أو باسم المستفيد — هذا يتطلب التحقق من البنك "
        "(خدمة التحقق من اسم المستفيد عبر سريع / ساريع).",
        "المرجع: ISO 13616-1:2020، ISO 7064 (MOD 97-10)، سجل SWIFT للآيبان، ومقال medium.com/@matlabb «IBAN Structure and MOD-97 Validation Algorithm».",
    ]
    for i, t in enumerate(notes, start=21):
        ws.cell(row=i, column=1, value=t).font = F_NOTE
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.cell(row=i, column=1).alignment = WRAP
        ws.row_dimensions[i].height = 32


# The guide file the user supplied (saudi_bank_codes.xlsx), row by row: (type, bank, its IBAN code, its SWIFT code,
# verified IBAN code or "غير مؤكد", verified SWIFT or "غير مؤكد", note).
UNVERIFIED = "غير مؤكد"
GUIDE = [
    ("محلي", "مصرف الراجحي", "80", "RJHISARI", "80", "RJHISARI", ""),
    ("محلي", "البنك الأهلي السعودي (SNB)", "10", "NCBKSARI", "10", "NCBKSAJE", "رمز SWIFT للأهلي ينتهي بـ JE وليس RI"),
    ("محلي", "بنك الرياض", "20", "RYADSARI", "20", "RIBLSARI", "رمز SWIFT لبنك الرياض هو RIBLSARI"),
    ("محلي", "مصرف الإنماء", "05", "INMASARI", "05", "INMASARI", ""),
    ("محلي", "بنك البلاد", "15", "BLADSARI", "15", "ALBISARI", "رمز SWIFT لبنك البلاد هو ALBISARI"),
    ("محلي", "البنك العربي الوطني (ANB)", "30", "ARNBSARI", "30", "ARNBSARI", ""),
    ("محلي", "البنك السعودي الأول (SAB)", "40 / 45", "SABBSARI", "45", "SABBSARI",
     "40 كان رمز سامبا (اندمج في الأهلي 2021) وليس ساب؛ رمز ساب 45 فقط (و50 للبنك الأول المدمج فيه)"),
    ("محلي", "البنك السعودي الفرنسي", "55", "BSFRSARI", "55", "BSFRSARI", ""),
    ("محلي", "بنك الجزيرة", "60", "BJAZSARI", "60", "BJAZSAJE", "رمز SWIFT لبنك الجزيرة BJAZSAJE (جدة)"),
    ("محلي", "البنك السعودي للاستثمار (SAIB)", "65", "SBVCSARI", "65", "SIBCSARI", "رمز SWIFT للبنك السعودي للاستثمار SIBCSARI"),
    ("رقمي", "بنك إس تي سي (STC Bank)", "76", "STCBSARI", "78", "STCJSARI",
     "76 رمز بنك مسقط؛ رمز STC Bank هو 78 (آيبان منشور يجتاز MOD-97، مصدر واحد) ورمز SWIFT هو STCJSARI"),
    ("رقمي", "بنك D360", "36", "D360SARI", "36", "DBAKSARI",
     "الرمز 36 مؤكد بآيبان فعلي قدّمه المستخدم يجتاز MOD-97؛ رمز SWIFT الصحيح DBAKSARI (أو DBAKSARIXXX بصيغة 11 خانة للمركز الرئيسي)"),
    ("أجنبي", "بنك أبوظبي الأول (FAB)", "73", "FABMSARI", "73", "FABMSARI",
     "الرمز 73 مؤكد بآيبان فعلي قدّمه المستخدم يجتاز MOD-97؛ FABMSARI ورد في مصدر واحد"),
    ("أجنبي", "بنك الكويت الوطني (NBK)", "75", "NBOKSAJE", "75", "NBOKSAJE", ""),
    ("أجنبي", "بنك الكويت الوطني (NBK) - إصدارات أخرى", "95", "NBOKSAJE", "75", "NBOKSAJE", "95 هو رمز بنك الإمارات دبي الوطني وليس NBK"),
    ("أجنبي", "بنك الإمارات دبي الوطني", "71", "EBILSARI", "95", "EBILSARI", "71 هو رمز بنك البحرين الوطني"),
    ("أجنبي", "بنك الخليج الدولي (GIB / meem)", "51", "GULFSARI", "90", UNVERIFIED, "GULFSARI ورد في الملف الاسترشادي فقط"),
    ("أجنبي", "بنك مسقط", "90", "BMAYSARI", "76", "BMUSSARI", "90 هو رمز بنك الخليج الدولي؛ SWIFT بنك مسقط BMUSSARI (مصدر واحد)"),
    ("أجنبي", "بنك البحرين الوطني (NBB)", "85", "NBBHSARI", "71", "NBOBSARI", "85 هو رمز بي إن بي باريبا؛ SWIFT بنك البحرين الوطني NBOBSARI (مصدر واحد)"),
    ("أجنبي", "جي بي مورغان تشيس", "78", "CHESSARI", "86", UNVERIFIED, "78 هو رمز STC Bank؛ رمز SWIFT المتداول CHASSARI وليس CHESSARI لكنه غير مؤكد"),
    ("أجنبي", "بي إن بي باريبا", "79", "BNPASARI", "85", UNVERIFIED, ""),
    ("أجنبي", "بنك الصين (Bank of China)", "56", "BKCHSARI", UNVERIFIED, UNVERIFIED, "لم يُعثر على مصدر موثوق"),
    ("أجنبي", "البنك الصناعي والتجاري الصيني (ICBC)", "53", "ICBKSARI", UNVERIFIED, UNVERIFIED, "لم يُعثر على مصدر موثوق"),
]


def build_audit(wb):
    ws = wb.create_sheet("تدقيق الملف الاسترشادي")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "تدقيق الملف الاسترشادي saudi_bank_codes.xlsx مقابل المصادر"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("الأعمدة C و D كما وردت في الملف؛ E و F القيمة المتحقق منها؛ G يبيّن لمن يعود الرمز الوارد في الملف فعلاً حسب ورقة «رموز البنوك». "
                "«غير مؤكد» تعني أنه لم يُعثر على مصدر موثوق، لا أن القيمة خاطئة بالضرورة.")
    ws["A2"].font = F_NOTE
    ws.merge_cells("A2:J2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 30
    header(ws, 4, ["النوع", "اسم البنك", "رمز الآيبان في الملف", "SWIFT في الملف", "رمز الآيبان المتحقق منه", "SWIFT المتحقق منه",
                   "الرمز في الملف يعود فعلاً إلى", "حكم رمز الآيبان", "حكم SWIFT", "ملاحظة"],
           [9, 34, 12, 14, 12, 14, 34, 16, 16, 52])
    for i, (kind, name, code, swift, vcode, vswift, note) in enumerate(GUIDE, start=5):
        for col, v in enumerate((kind, name, code, swift, vcode, vswift), start=1):
            ws.cell(row=i, column=col, value=v).number_format = "@"
        ws[f"G{i}"] = f'=IFERROR(INDEX({NAME_RANGE},MATCH(C{i},{BANK_RANGE},0)),"غير مسجل في الجدول")'
        ws[f"H{i}"] = f'=IF(E{i}="{UNVERIFIED}","? غير مؤكد",IF(C{i}=E{i},"✓ صحيح","✗ خطأ"))'
        ws[f"I{i}"] = f'=IF(F{i}="{UNVERIFIED}","? غير مؤكد",IF(D{i}=F{i},"✓ صحيح","✗ خطأ"))'
        ws[f"J{i}"] = note
        style_row(ws, i, 10)
        for col in (2, 7, 10):
            ws.cell(row=i, column=col).alignment = WRAP
        ws.row_dimensions[i].height = 30
    end = 4 + len(GUIDE)
    for rng in (f"H5:I{end}",):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(H5,1)="✓"'], fill=GREEN))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(H5,1)="✗"'], fill=RED))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(H5,1)="?"'], fill=FILL_INPUT))
    r = end + 2
    ws[f"A{r}"] = "الملخص"
    ws[f"A{r}"].font = F_BOLD
    for j, (label, f) in enumerate([("رموز آيبان صحيحة", f'=COUNTIF(H5:H{end},"✓*")'),
                                    ("رموز آيبان خاطئة", f'=COUNTIF(H5:H{end},"✗*")'),
                                    ("رموز آيبان غير مؤكدة", f'=COUNTIF(H5:H{end},"~?*")'),
                                    ("رموز SWIFT صحيحة", f'=COUNTIF(I5:I{end},"✓*")'),
                                    ("رموز SWIFT خاطئة", f'=COUNTIF(I5:I{end},"✗*")'),
                                    ("رموز SWIFT غير مؤكدة", f'=COUNTIF(I5:I{end},"~?*")')], start=r + 1):
        ws[f"B{j}"], ws[f"C{j}"] = label, f
        style_row(ws, j, 3)
    ws[f"A{r + 8}"] = ("بنوك ناقصة من الملف: دويتشه بنك (81)، البنك الوطني الباكستاني (82)، بنك الدولة الهندي (83)، بنك زراعات (84)، "
                       "ورموز البنوك المدمجة سامبا (40) والبنك الأول (50) التي قد تظهر في آيبانات قديمة.")
    ws[f"A{r + 8}"].font = Font(name=FONT, size=10, bold=True, color="C00000")
    ws.merge_cells(f"A{r + 8}:J{r + 8}")
    ws[f"A{r + 8}"].alignment = WRAP
    ws.row_dimensions[r + 8].height = 30
    ws.freeze_panes = "C5"


def main():
    wb = Workbook()
    build_structure(wb)
    build_banks(wb)
    build_validator(wb)
    build_algorithm(wb)
    build_audit(wb)
    wb.active = 2
    PACKAGE.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    write_json()
    print(OUT)


STATUS = {"نشط": "active", "بنك رقمي": "digital", "فرع بنك أجنبي": "foreign_branch", "مدمج (تاريخي)": "legacy_merged"}
CONFIDENCE = {CONFIRMED: "confirmed", ONE_SOURCE: "single_source", UNCONFIRMED: "unconfirmed"}


def write_json():
    banks = []
    for code, ar, en, status, _note in BANKS:
        swift, level = SWIFT[code]
        banks.append({"code": code, "name_ar": ar, "name_en": en, "status": STATUS[status], "swift": swift,
                      "swift_confidence": next(v for k, v in CONFIDENCE.items() if level.startswith(k))})
    doc = {"country": "SA", "iban_length": 24, "bban_format": "2!n18!c", "updated": "2026-10-30",
           "note": "IBAN bank identifiers assigned by SAMA, compiled from public sources; see README.md for sources and confidence.",
           "banks": banks}
    (PACKAGE / "saudi_banks.json").write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
